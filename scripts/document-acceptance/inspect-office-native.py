"""Reopen Excel's native fixed-format output and enforce printable context."""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import fitz
from openpyxl import load_workbook


root = Path(sys.argv[1]).resolve()
root.relative_to(Path("output").resolve())

expectations = {
    "class-roster-types-native.pdf": {
        "all": ("Sıra", "Adı soyadı", "000012", "10.09.2020"),
        "none": (),
        "every_page": ("Sıra", "Adı soyadı", "Sınıf:"),
    },
    "class-roster-full-native.pdf": {
        "all": ("Sıra", "Adı soyadı", "Kurgu İpek"),
        "none": (),
        "every_page": ("Sıra", "Adı soyadı", "Sınıf:"),
    },
    "class-roster-compact-25-native.pdf": {
        "all": ("Sıra", "Öğrenci"),
        "normalized_all": (
            "+90 (532) 111 22 33 / +90 (258) 444 55 66 / dahili 777",
        ),
        "none": ("SADECE_TAM_VERI_ADRES_987654", "SADECE_TAM_VERI_MESLEK_2468"),
        "every_page": ("Sıra", "Öğrenci", "Sınıf:"),
        "worksheet": "Kısa iletişim baskısı",
        "name_column": 3,
        "first_data_row": 6,
        "max_pages": 2,
        "minimum_body_font_size": 6.5,
    },
    "class-roster-compact-edit-probe-native.pdf": {
        "all": ("Sıra", "Öğrenci", "0532 777 66 55"),
        "normalized_all": (
            "0532 111 22 33",
            '=HYPERLINK("https://example.invalid","Kurgu")',
        ),
        "none": ("SADECE_TAM_VERI_ADRES_987654", "SADECE_TAM_VERI_MESLEK_2468"),
        "every_page": ("Sıra", "Öğrenci", "Sınıf:"),
        "worksheet": "Kısa iletişim baskısı",
        "name_column": 3,
        "first_data_row": 6,
        "max_pages": 2,
        "minimum_body_font_size": 6.5,
    },
    "growth-class-40-native.pdf": {
        "all": (
            "Öğrenci",
            "Kurgu Ada",
            "Kurgu Çok Uzun İsimli",
            "Okul Öncesi Öğretmeni: Kurgu Öğretmen",
            "İmza:",
        ),
        "none": (),
        "every_page": (
            "Sıra",
            "Öğrenci",
            "Dönem",
            "Kurgu Anaokulu",
            "Kurgu Deniz Sınıfı",
            "2026–2027 Kurgu Eğitim Yılı",
        ),
    },
    "growth-individual-native.pdf": {
        "all": ("Öğrenci", "Kurgu Ada", "Okul Öncesi Öğretmeni: Kurgu Öğretmen", "İmza:"),
        "none": ("Kurgu Bora",),
        "every_page": (
            "Sıra",
            "Öğrenci",
            "Dönem",
            "Kurgu Anaokulu",
            "Kurgu Deniz Sınıfı",
            "2026–2027 Kurgu Eğitim Yılı",
        ),
    },
    "growth-class-september-native.pdf": {
        "all": (
            "Öğrenci",
            "Kurgu Ada",
            "Eylül 2026",
            "Okul Öncesi Öğretmeni: Kurgu Öğretmen",
            "İmza:",
        ),
        "none": ("Aralık", "Mart", "Haziran", "123,4 cm", "20,125 kg"),
        "every_page": (
            "Sıra",
            "Öğrenci",
            "Dönem",
            "Eylül 2026",
            "Kurgu Anaokulu",
            "Kurgu Deniz Sınıfı",
            "2026–2027 Kurgu Eğitim Yılı",
        ),
    },
}

receipt: dict[str, object] = {"files": [], "failures": [], "passed": False}
failures: list[str] = receipt["failures"]  # type: ignore[assignment]

for name, expected in expectations.items():
    path = root / name
    if not path.is_file():
        failures.append(f"{name}: native Excel PDF bulunamadı")
        continue
    pages: list[dict[str, object]] = []
    texts: list[str] = []
    with fitz.open(path) as document:
        maximum_pages = expected.get("max_pages")
        if maximum_pages is not None and document.page_count > maximum_pages:
            failures.append(
                f"{name}: {document.page_count} sayfa üretti; üst sınır {maximum_pages}"
            )
        for index, page in enumerate(document):
            text = page.get_text()
            texts.append(text)
            spans = [
                span
                for block in page.get_text("dict")["blocks"]
                if "lines" in block
                for line in block["lines"]
                for span in line["spans"]
                if span["text"].strip()
            ]
            outside = [
                span["bbox"]
                for span in spans
                if span["bbox"][0] < -0.5
                or span["bbox"][1] < -0.5
                or span["bbox"][2] > page.rect.width + 0.5
                or span["bbox"][3] > page.rect.height + 0.5
            ]
            body_font_sizes = [
                float(span["size"])
                for span in spans
                if span["bbox"][3] < page.rect.height - 24
            ]
            minimum_body_font_size = min(body_font_sizes, default=0.0)
            required_body_font_size = expected.get("minimum_body_font_size")
            if required_body_font_size is not None and minimum_body_font_size < required_body_font_size:
                failures.append(
                    f"{name} sayfa {index + 1}: gövde yazısı {minimum_body_font_size:.2f} pt; "
                    f"alt sınır {required_body_font_size:.2f} pt"
                )
            missing_context = [value for value in expected["every_page"] if value not in text]
            if outside:
                failures.append(f"{name} sayfa {index + 1}: metin sayfa sınırı dışında")
            if not text.strip():
                failures.append(f"{name} sayfa {index + 1}: aranabilir metin yok")
            if missing_context:
                failures.append(
                    f"{name} sayfa {index + 1}: yinelenen bağlam eksik: {', '.join(missing_context)}"
                )
            pages.append(
                {
                    "page": index + 1,
                    "textCharacters": len(text),
                    "outsideCount": len(outside),
                    "missingContext": missing_context,
                    "minimumBodyFontSize": round(minimum_body_font_size, 3),
                }
            )
        metadata = document.metadata or {}
    all_text = "\n".join(texts)
    # Compare complete printable names from an independent XLSX reader, including
    # names that wrap over many lines; checking only their prefix misses clipping.
    is_growth = name.startswith("growth-")
    workbook = load_workbook(root / name.replace("-native.pdf", ".xlsx"), read_only=True, data_only=True)
    worksheet = workbook[
        expected.get("worksheet", "Baskı çizelgesi" if is_growth else "Sınıf listesi")
    ]
    name_column = int(expected.get("name_column", 2 if is_growth else 3))
    first_data_row = int(expected.get("first_data_row", 5 if is_growth else 8))
    printable_names = {
        str(row[name_column - 1])
        for row in worksheet.iter_rows(min_row=first_data_row, values_only=True)
        if len(row) >= name_column and row[name_column - 1] and isinstance(row[0], (int, float))
    }
    workbook.close()
    normalized_text = re.sub(r"\s+", "", all_text)
    missing_names = [value for value in printable_names if re.sub(r"\s+", "", value) not in normalized_text]
    if missing_names:
        failures.append(f"{name}: {len(missing_names)} tam öğrenci adı baskı metninde eksik")
    for value in expected["all"]:
        if value not in all_text:
            failures.append(f"{name}: beklenen metin yok: {value}")
    for value in expected.get("normalized_all", ()):
        if re.sub(r"\s+", "", value) not in normalized_text:
            failures.append(f"{name}: beklenen tam metin yok: {value}")
    for value in expected["none"]:
        if value in all_text:
            failures.append(f"{name}: kapsam dışı metin sızdı: {value}")
    receipt["files"].append(  # type: ignore[union-attr]
        {
            "name": name,
            "bytes": path.stat().st_size,
            "pages": pages,
            "metadataTitle": metadata.get("title", ""),
            "completeNameCount": len(printable_names),
            "missingCompleteNameCount": len(missing_names),
        }
    )

receipt["passed"] = not failures and len(receipt["files"]) == len(expectations)  # type: ignore[arg-type]
(root / "native-excel-pdf-readback.json").write_text(
    json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
)
print(json.dumps(receipt, ensure_ascii=False, indent=2))
raise SystemExit(0 if receipt["passed"] else 1)
