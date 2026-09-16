from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import json
import pypdfium2 as pdfium
from pypdf import PdfReader
from openpyxl import load_workbook

folder = Path('output/teacher-report-center/final')
proof = json.loads((folder / 'proof.json').read_text(encoding='utf-8'))
word = next(folder.glob('*week*.docx'))
with ZipFile(word) as z:
    root = ET.fromstring(z.read('word/document.xml'))
    ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    size = root.find('.//' + ns + 'pgSz')
    assert size.attrib[ns + 'orient'] == 'landscape'
    assert int(size.attrib[ns + 'w']) > int(size.attrib[ns + 'h'])
    assert proof['teacher'] in z.read('word/footer1.xml').decode('utf-8')
with ZipFile(next(folder.glob('*year*.docx'))) as z:
    word_text = ''.join(ET.fromstring(z.read('word/document.xml')).itertext())
    assert proof['summary'] in word_text
    assert proof['closure'][0]['text'] in word_text
book = load_workbook(next(folder.glob('*.xlsx')), data_only=True)
assert proof['summary'] in [r[0] for r in book['Yönetici özeti'].iter_rows(values_only=True)]
assert any(proof['closure'][0]['text'] in str(r[5]) for r in book['Veri'].iter_rows(values_only=True))
pdf_path = next(folder.glob('*year*.pdf'))
reader = PdfReader(pdf_path)
text = ' '.join(' '.join(p.extract_text().split()) for p in reader.pages)
assert proof['summary'] in text
assert proof['closure'][0]['text'] in text
pdf = pdfium.PdfDocument(pdf_path)
for i in range(len(pdf)):
    pdf[i].render(scale=1.2).to_pil().save(folder / f'annual-{i+1}.png')
(folder / 'layout-proof.json').write_text(json.dumps({'weeklyWordLandscape': True, 'teacherFooter': True, 'sameSummaryPdfWordExcel': True, 'actualNotePdfWordExcel': True, 'annualPdfPages': len(pdf)}, indent=2), encoding='utf-8')
print((folder / 'layout-proof.json').read_text())
