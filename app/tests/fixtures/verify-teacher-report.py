from pathlib import Path
import json
import pypdfium2 as pdfium
from pypdf import PdfReader
from openpyxl import load_workbook

folder=Path('output/teacher-report-center')
proof={}
for name in ['report','long-report','report-word','long-report-word','report-excel','long-report-excel']:
    path=folder/f'{name}.pdf'
    reader=PdfReader(path)
    text='\n'.join(page.extract_text() or '' for page in reader.pages)
    text=' '.join(text.split())
    assert 'Kurgu Çocuk 2' not in text, name
    assert not any(error in text for error in ['#REF!','#VALUE!','#N/A']),name
    if name.startswith('long-report'):
        assert text.count('korunmalı.')==100,(name,text.count('korunmalı.'))
    pdf=pdfium.PdfDocument(path)
    pdf[0].render(scale=1.2).to_pil().save(folder/f'{name}.png')
    if len(pdf)>1:pdf[len(pdf)-1].render(scale=1.2).to_pil().save(folder/f'{name}-last.png')
    proof[name]={'pages':len(reader.pages),'long_phrase_count':text.count('Uzun kaynak eksiksiz korunmalı.')}
book=load_workbook(folder/'long-report.xlsx',data_only=True)
assert book.sheetnames==['Veri','Yazdırma','Açıklamalar']
assert [r[5].value for r in book['Veri'].iter_rows(min_row=2)]==[r[3].value for r in book['Yazdırma'].iter_rows(min_row=2)], 'Native linked values differ'
assert ''.join(str(row[5].value or '') for row in book['Veri'].iter_rows(min_row=2)).count('Uzun kaynak eksiksiz korunmalı.')==100
proof['native_excel']={'rows':book['Yazdırma'].max_row,'all_values_recalculated':True}
(folder/'native-proof.json').write_text(json.dumps(proof,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(proof,ensure_ascii=False))

